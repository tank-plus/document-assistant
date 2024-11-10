import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

# 创建一个新的工作簿和工作表
wb = openpyxl.Workbook()
ws = wb.active

# 写入多维数据：Category, Sales, Prices
data = [
    ['Category', 'Sales', 'Prices'],
    ['A', 30, 10],
    ['B', 70, 20],
    ['C', 50, 15],
    ['D', 90, 25]
]

for row in data:
    ws.append(row)

# 生成柱状图 (根据Sales)
bar_chart = BarChart()
bar_chart.title = "Sales by Category"
bar_chart.x_axis.title = "Category"
bar_chart.y_axis.title = "Sales"

# 引用Sales数据区域
categories = Reference(ws, min_col=1, min_row=2, max_row=5)
sales_values = Reference(ws, min_col=2, min_row=1, max_row=5)
bar_chart.add_data(sales_values, titles_from_data=True)
bar_chart.set_categories(categories)

# 将柱状图添加到工作表
ws.add_chart(bar_chart, "E5")

# 生成饼图 (根据Sales)
pie_chart = PieChart()
pie_chart.title = "Sales Distribution"

# 饼图使用Sales数据
pie_chart.add_data(sales_values, titles_from_data=True)
pie_chart.set_categories(categories)

# 将饼图添加到工作表
ws.add_chart(pie_chart, "E20")

# 生成折线图 (根据Sales和Prices)
line_chart = LineChart()
line_chart.title = "Sales and Prices Trend"
line_chart.x_axis.title = "Category"
line_chart.y_axis.title = "Value"

# 引用Prices数据
prices_values = Reference(ws, min_col=3, min_row=1, max_row=5)
line_chart.add_data(sales_values, titles_from_data=True)
line_chart.add_data(prices_values, titles_from_data=True)
line_chart.set_categories(categories)

# 将折线图添加到工作表
ws.add_chart(line_chart, "E35")

# 生成Prices的柱状图
bar_chart_prices = BarChart()
bar_chart_prices.title = "Prices by Category"
bar_chart_prices.x_axis.title = "Category"
bar_chart_prices.y_axis.title = "Prices"

# 引用Prices数据
bar_chart_prices.add_data(prices_values, titles_from_data=True)
bar_chart_prices.set_categories(categories)

# 将Prices的柱状图添加到工作表
ws.add_chart(bar_chart_prices, "E50")

# 保存Excel文件
wb.save("chart_example_with_sales_and_prices.xlsx")
