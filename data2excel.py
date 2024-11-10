import pandas as pd
import mysql.connector
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

# Step 1: 从MySQL读取数据
# 替换为你自己的MySQL数据库连接信息
connection = mysql.connector.connect(
    host='36.111.148.127',
    port='13306',
    user='root',
    password='passw0rd',
    database='product'
)

# SQL查询：假设商品数据表包含 name, price, sales 字段
query = "SELECT name, price, sales FROM product"
df = pd.read_sql(query, connection)

# 关闭数据库连接
connection.close()

# Step 2: 过滤价格大于100的数据
filtered_df = df[df['price'] > 100][['name', 'price', 'sales']]

# Step 3: 将数据写入Excel文件
excel_file = 'filtered_products.xlsx'
filtered_df.to_excel(excel_file, index=False)

# Step 4: 使用openpyxl加载Excel文件并创建柱状图和折线图
wb = load_workbook(excel_file)
ws = wb.active

# 创建柱状图 (显示销量)
bar_chart = BarChart()
bar_chart.title = "Sales by Product"
bar_chart.x_axis.title = "Product"
bar_chart.y_axis.title = "Sales"

# 引用产品名称和销量数据
categories = Reference(ws, min_col=1, min_row=2, max_row=len(filtered_df) + 1)
sales_values = Reference(ws, min_col=3, min_row=1, max_row=len(filtered_df) + 1)
bar_chart.add_data(sales_values, titles_from_data=True)
bar_chart.set_categories(categories)

# 将柱状图添加到工作表
ws.add_chart(bar_chart, "E5")

# 创建折线图 (显示价格和销量的趋势)
line_chart = LineChart()
line_chart.title = "Price and Sales Trend"
line_chart.x_axis.title = "Product"
line_chart.y_axis.title = "Values"

# 引用价格数据
price_values = Reference(ws, min_col=2, min_row=1, max_row=len(filtered_df) + 1)
line_chart.add_data(sales_values, titles_from_data=True)
line_chart.add_data(price_values, titles_from_data=True)
line_chart.set_categories(categories)

# 将折线图添加到工作表
ws.add_chart(line_chart, "E20")

# 保存带有图表的Excel文件
wb.save(excel_file)

print(f"Excel file saved as {excel_file}")
